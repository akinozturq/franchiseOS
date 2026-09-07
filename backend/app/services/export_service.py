from io import BytesIO
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.app.schemas.reconciliation import ReconciliationResponse

def generate_reconciliation_excel(report: ReconciliationResponse) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Mutabakat {report.month:02d}-{report.year}"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name="Calibri", size=16, bold=True, color="1F2937")
    font_section = Font(name="Calibri", size=12, bold=True, color="111827")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="1F2937")
    font_regular = Font(name="Calibri", size=11, color="1F2937")
    
    fill_header_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header_teal = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    fill_gray_row = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )
    
    # 1. Main Header
    ws.merge_cells("A1:E1")
    ws["A1"] = f"{report.branch_name} — DÖNEMLİK MUTABAKAT VE KOMİSYON RAPORU"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")
    
    ws["A2"] = f"Dönem: {report.month:02d}/{report.year} | Tahsilat Yapan Taraf: {report.invoice_summary.collector_party.value}"
    ws["A2"].font = font_regular
    
    # 2. General Summary Section
    ws["A4"] = "1. GENEL MUTABAKAT VE GELİR PAYLAŞIMI ÖZETİ"
    ws["A4"].font = font_section
    
    summary_headers = ["Gösterge / Kalem", "Değer", "Oran / Açıklama"]
    for col_idx, h in enumerate(summary_headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    summary_data = [
        ("Toplam İşlem Adedi", f"{report.total_transactions} adet", "-"),
        ("Toplam Dönem Cirosu (KDV Hariç)", float(report.total_turnover_excl_vat), "%100,00"),
        ("Uygulanan Dilim Oranı (Bayi)", f"%{float(report.applied_rate_percentage):.2f}", report.rate_explanation),
        ("Bayi Payı (KDV Hariç)", float(report.bayi_share_excl_vat), f"%{float(report.bayi_share_rate * 100):.2f}"),
        ("Franchisor (Merkez) Payı (KDV Hariç)", float(report.franchisor_share_excl_vat), f"%{float(report.franchisor_share_rate * 100):.2f}"),
    ]
    
    for r_idx, row in enumerate(summary_data, start=6):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            if isinstance(val, float):
                cell.number_format = "#,##0.00 TL"
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 1:
                cell.font = font_bold
                
    # 3. Invoicing Section
    inv_start_row = 12
    ws.cell(row=inv_start_row, column=1, value="2. DÖNEM SONU KESİLMESİ GEREKEN FATURA BİLGİSİ").font = font_section
    
    inv = report.invoice_summary
    inv_headers = ["Fatura Kalemi", "Açıklama / Tutar"]
    for col_idx, h in enumerate(inv_headers, start=1):
        cell = ws.cell(row=inv_start_row + 1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_teal
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        
    invoice_rows = [
        ("Faturayı Kesen Taraf (Satıcı)", inv.issuer),
        ("Faturayı Alan Taraf (Alıcı)", inv.recipient),
        ("Fatura Açıklaması", inv.description),
        ("Fatura Tutarı (KDV Hariç)", float(inv.amount_excl_vat)),
        (f"KDV Oranı (%{float(inv.vat_rate * 100):.0f})", float(inv.vat_amount)),
        ("GENEL TOPLAM (KDV DAHİL ÖDENECEK)", float(inv.total_amount_incl_vat)),
    ]
    
    for r_idx, (k, v) in enumerate(invoice_rows, start=inv_start_row + 2):
        cell_k = ws.cell(row=r_idx, column=1, value=k)
        cell_v = ws.cell(row=r_idx, column=2, value=v)
        cell_k.border = thin_border
        cell_v.border = thin_border
        cell_k.font = font_bold
        
        if isinstance(v, float):
            cell_v.number_format = "#,##0.00 TL"
            cell_v.alignment = Alignment(horizontal="right")
            cell_v.font = font_regular
        else:
            cell_v.font = font_regular
            
        if "GENEL TOPLAM" in k:
            cell_k.fill = fill_highlight
            cell_v.fill = fill_highlight
            cell_v.font = font_bold

    # 4. Department Breakdown Section
    dept_start_row = inv_start_row + len(invoice_rows) + 3
    ws.cell(row=dept_start_row, column=1, value="3. DEPARTMAN BAZLI CİRO KIRILIMI").font = font_section
    
    dept_headers = ["Departman Adı", "İşlem Adedi", "KDV Hariç Ciro (TL)", "Toplam Payı (%)"]
    for col_idx, h in enumerate(dept_headers, start=1):
        cell = ws.cell(row=dept_start_row + 1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        
    cur_row = dept_start_row + 2
    for dept in report.department_breakdown:
        ws.cell(row=cur_row, column=1, value=dept.department_name).border = thin_border
        cell_cnt = ws.cell(row=cur_row, column=2, value=dept.transaction_count)
        cell_cnt.alignment = Alignment(horizontal="center")
        cell_cnt.border = thin_border
        
        cell_amt = ws.cell(row=cur_row, column=3, value=float(dept.total_turnover_excl_vat))
        cell_amt.number_format = "#,##0.00 TL"
        cell_amt.border = thin_border
        
        cell_pct = ws.cell(row=cur_row, column=4, value=f"%{float(dept.percentage_of_total):.2f}")
        cell_pct.alignment = Alignment(horizontal="right")
        cell_pct.border = thin_border
        cur_row += 1
        
    # Department Total Row
    cell_tot_lbl = ws.cell(row=cur_row, column=1, value="GENEL TOPLAM")
    cell_tot_lbl.font = font_bold
    cell_tot_lbl.fill = fill_gray_row
    cell_tot_lbl.border = thin_border
    
    cell_tot_cnt = ws.cell(row=cur_row, column=2, value=report.total_transactions)
    cell_tot_cnt.font = font_bold
    cell_tot_cnt.fill = fill_gray_row
    cell_tot_cnt.alignment = Alignment(horizontal="center")
    cell_tot_cnt.border = thin_border
    
    cell_tot_amt = ws.cell(row=cur_row, column=3, value=float(report.total_turnover_excl_vat))
    cell_tot_amt.font = font_bold
    cell_tot_amt.fill = fill_gray_row
    cell_tot_amt.number_format = "#,##0.00 TL"
    cell_tot_amt.border = thin_border
    
    cell_tot_pct = ws.cell(row=cur_row, column=4, value="%100,00")
    cell_tot_pct.font = font_bold
    cell_tot_pct.fill = fill_gray_row
    cell_tot_pct.alignment = Alignment(horizontal="right")
    cell_tot_pct.border = thin_border

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 16)
        
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 45

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_bonus_excel(report, branch_name: str = "Merkez Bayi") -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Personel Prim {report.period}"

    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Calibri", size=15, bold=True, color="1F2937")
    font_section = Font(name="Calibri", size=11, bold=True, color="111827")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True, color="1F2937")
    font_regular = Font(name="Calibri", size=10, color="1F2937")

    fill_header_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_gray_row = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    # 1. Header
    ws.merge_cells("A1:K1")
    ws["A1"] = f"{branch_name} — PERSONEL DÖNEMLİK PRİM VE HAKEDİŞ RAPORU"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")

    ws["A2"] = f"Dönem: {report.period} | Toplam Personel: {report.total_employees} | Toplam Dağıtılan Prim: {float(report.grand_total_bonus):,.2f} TL"
    ws["A2"].font = font_regular

    # 2. Table Headers
    headers = [
        "Personel Adı",
        "Rol",
        "Departman",
        "Ciro Kaynağı",
        "İşlem Ad.",
        "Standart Ciro (TL)",
        "Dilim Oranı",
        "Standart Prim (TL)",
        "İstisnalı Ciro (TL)",
        "İstisnalı Prim (TL)",
        "Toplam Ciro (TL)",
        "TOPLAM PRİM (TL)",
        "Efektif Oran",
        "Hesaplama Açıklaması"
    ]

    header_row = 4
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[header_row].height = 28

    # Turnover source translation
    source_labels = {
        "kendi_islemleri": "Kendi İşlemleri",
        "kendi_departmani": "Kendi Departmanı",
        "tum_bayi": "Tüm Bayi"
    }

    cur_row = 5
    tot_std_turnover = Decimal("0.00")
    tot_std_bonus = Decimal("0.00")
    tot_ovr_turnover = Decimal("0.00")
    tot_ovr_bonus = Decimal("0.00")
    tot_turnover = Decimal("0.00")
    tot_bonus = Decimal("0.00")

    for item in report.items:
        tot_std_turnover += item.standard_turnover_excl_vat
        tot_std_bonus += item.standard_bonus
        tot_ovr_turnover += item.override_turnover_excl_vat
        tot_ovr_bonus += item.override_bonus
        tot_turnover += item.total_turnover_excl_vat
        tot_bonus += item.total_bonus

        row_vals = [
            (item.employee_name, "left", False, None),
            (item.role_name, "left", False, None),
            (item.department_name or "Tüm Bayi", "left", False, None),
            (source_labels.get(item.turnover_source, item.turnover_source), "center", False, None),
            (item.transaction_count, "center", False, "#,##0"),
            (float(item.standard_turnover_excl_vat), "right", False, "#,##0.00 TL"),
            (f"%{float(item.standard_tier_rate * 100):.1f}", "right", False, None),
            (float(item.standard_bonus), "right", False, "#,##0.00 TL"),
            (float(item.override_turnover_excl_vat), "right", False, "#,##0.00 TL"),
            (float(item.override_bonus), "right", False, "#,##0.00 TL"),
            (float(item.total_turnover_excl_vat), "right", False, "#,##0.00 TL"),
            (float(item.total_bonus), "right", True, "#,##0.00 TL"),
            (f"%{float(item.effective_bonus_rate * 100):.2f}", "right", False, None),
            (item.tier_explanation, "left", False, None),
        ]

        for c_idx, (val, align, is_total_col, num_fmt) in enumerate(row_vals, start=1):
            cell = ws.cell(row=cur_row, column=c_idx, value=val)
            cell.font = font_bold if is_total_col else font_regular
            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.border = thin_border
            if num_fmt and isinstance(val, (int, float)):
                cell.number_format = num_fmt
            if is_total_col:
                cell.fill = fill_highlight

        cur_row += 1

    # Grand Total Row
    tot_row = cur_row
    ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=4)
    cell_lbl = ws.cell(row=tot_row, column=1, value="GENEL TOPLAM")
    cell_lbl.font = font_bold
    cell_lbl.fill = fill_gray_row
    cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
    
    for c in range(1, 5):
        ws.cell(row=tot_row, column=c).border = thin_border
        ws.cell(row=tot_row, column=c).fill = fill_gray_row

    total_cells = [
        (5, sum(i.transaction_count for i in report.items), "#,##0", "center"),
        (6, float(tot_std_turnover), "#,##0.00 TL", "right"),
        (7, "-", None, "center"),
        (8, float(tot_std_bonus), "#,##0.00 TL", "right"),
        (9, float(tot_ovr_turnover), "#,##0.00 TL", "right"),
        (10, float(tot_ovr_bonus), "#,##0.00 TL", "right"),
        (11, float(tot_turnover), "#,##0.00 TL", "right"),
        (12, float(tot_bonus), "#,##0.00 TL", "right"),
        (13, f"%{float(tot_bonus / tot_turnover * 100):.2f}" if tot_turnover > 0 else "%0.00", None, "right"),
        (14, f"{report.total_employees} personelin toplam hak edişi", None, "left")
    ]

    for col_idx, val, num_fmt, align in total_cells:
        cell = ws.cell(row=tot_row, column=col_idx, value=val)
        cell.font = font_bold
        cell.fill = fill_gray_row if col_idx != 12 else fill_highlight
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.border = thin_border
        if num_fmt and isinstance(val, (int, float)):
            cell.number_format = num_fmt

    # Set column widths
    col_widths = {
        1: 22,  # Personel Adı
        2: 24,  # Rol
        3: 20,  # Departman
        4: 16,  # Ciro Kaynağı
        5: 10,  # İşlem Ad.
        6: 18,  # Standart Ciro
        7: 12,  # Dilim Oranı
        8: 18,  # Standart Prim
        9: 18,  # İstisnalı Ciro
        10: 18, # İstisnalı Prim
        11: 18, # Toplam Ciro
        12: 20, # Toplam Prim
        13: 13, # Efektif Oran
        14: 45  # Açıklama
    }

    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

